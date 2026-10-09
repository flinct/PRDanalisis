# Generate UAT prototype Sales -> md + xlsx from one source list.
# ponytail: skenario inline di sini = single source; regen dgn `python build_uat.py`.
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).parent
OUT = ROOT / "uat"
OUT.mkdir(exist_ok=True)

META = {
    "Judul": "UAT Prototype Modul Sales (SatuInbox)",
    "Apps Name": "SatuInbox — Sales Module (prototype)",
    "Version": "prototype 1.0",
    "Date": "2026-09-29",
    "PM": "Dany Christian",
    "QA Name": "Dany Christian",
    "TechLead": "Naftal Yunior",
    "Scope": "sales-clone/ (leads, visits, settings pipeline) — walkthrough manual di browser (file://)",
}

COLS = ["No", "Area", "Skenario", "Langkah", "Expected Result", "Pass/Fail", "Catatan"]

# (Area, Skenario, Langkah, Expected)
CASES = [
 ("Navigasi", "Landing default = Leads",
  "Buka index.html sales-clone tanpa hash.",
  "Halaman langsung ke Pipeline Leads (bukan blank / conversation). Rail hanya menampilkan Prospek."),
 ("Navigasi", "Visit tidak ada di main nav",
  "Perhatikan rail kiri.",
  "Tidak ada tombol Kunjungan/Visit di rail. Visit hanya diakses dari dalam Leads."),
 ("View", "Toggle Board / Table",
  "Klik toggle Board lalu Table.",
  "Board = kanban per stage; Table = list baris. Data lead sama di kedua view."),
 ("View", "Filter pencarian lead",
  "Ketik nama/company di kotak filter.",
  "List lead menyusut sesuai keyword secara live."),
 ("View", "Filter per stage (klik summary card)",
  "Klik salah satu KPI/stage card di ringkasan.",
  "List terfilter ke stage tsb; klik lagi = reset."),
 ("View", "Sort arah asc/desc",
  "Klik toggle arah sort.",
  "Urutan default descending; toggle membalik urutan."),
 ("View", "Follow-up filter",
  "Aktifkan filter follow-up.",
  "Hanya lead yang punya nextAction jatuh tempo yang tampil."),
 ("View", "Pipeline value breakdown",
  "Klik toggle breakdown nilai pipeline.",
  "Muncul rincian total nilai per stage (termasuk Won/Lost value KPI)."),
 ("Board", "Drag lead antar stage",
  "Board view: drag kartu lead ke kolom stage lain (non-terminal).",
  "Lead pindah stage; toast konfirmasi; data persist setelah refresh."),
 ("Detail", "Buka detail lead",
  "Klik satu kartu/baris lead.",
  "Modal detail terbuka: summary, stepper status, kontak, aktivitas, visit."),
 ("Detail", "Inline-edit field",
  "Edit field (mis. nilai/assignee) di detail lalu blur.",
  "Nilai tersimpan otomatis (blur-save); toast konfirmasi; persist."),
 ("Detail", "Tambah komentar/aktivitas",
  "Isi form aktivitas lalu simpan.",
  "Aktivitas baru muncul di timeline; toast konfirmasi."),
 ("Status Guard", "Transisi forward hanya next stage",
  "Buka popover status di detail, coba pindah maju.",
  "Hanya next stage yang bisa dipilih untuk maju; lompat jauh diblok."),
 ("Status Guard", "Won/Lost butuh alasan",
  "Pilih status terminal (Won/Lost).",
  "Wajib isi reason sebelum konfirmasi; tanpa reason tidak bisa lanjut."),
 ("Status Guard", "Terminal tidak bisa balik",
  "Lead yang sudah Won/Lost, coba ubah status.",
  "Status terminal terkunci / tidak bisa kembali ke stage aktif."),
 ("Create", "Toggle varian A/B/C",
  "Buka Create Lead, klik toggle A, B, C.",
  "Body form berubah (A detail-style, B ringkas, C sederhana); isian terjaga saat ganti varian."),
 ("Create", "Stepper status varian A (default new)",
  "Varian A: perhatikan stepper status, klik segmen.",
  "Default stage pertama non-terminal (new). Klik segmen mengisi bar sampai segmen itu (done+current). Won/Lost tak bisa dipilih di create."),
 ("Create", "Validasi Judul Lead min 3 karakter",
  "Submit dengan Judul Lead kosong / <3 karakter.",
  "Error 'Judul Lead minimal 3 karakter'; submit ditolak."),
 ("Create", "Validasi format email",
  "Isi email tidak valid lalu submit.",
  "Error 'Format email tidak valid'; submit ditolak."),
 ("Create", "Nilai negatif di-clamp",
  "Isi Nilai negatif lalu submit.",
  "Nilai tersimpan minimal 0 (tidak negatif)."),
 ("Create", "Autocomplete kontak",
  "Ketik di Nama PIC, pilih saran kontak.",
  "Field telepon/email terisi otomatis dari kontak; lead ter-link contactId."),
 ("Create", "Label general (Judul Lead / Nama PIC)",
  "Perhatikan label form.",
  "Label 'Judul Lead' (placeholder Perusahaan / perorangan) & 'Nama PIC', bukan 'Nama perusahaan'/'Nama Kontak'."),
 ("Create", "Simpan lead baru",
  "Isi form valid lalu simpan.",
  "Lead baru muncul di board/table pada stage yang dipilih; toast konfirmasi; persist."),
 ("Role", "Delete lead sesuai role",
  "Switch role (agent vs admin/spv) via role modal, cek aksi delete.",
  "Delete hanya tersedia untuk role berwenang; agent tidak bisa delete."),
 ("Visit", "Buka modal visit dari lead",
  "Di detail lead, buka visit / check-in.",
  "Modal visit terbuka; bisa buat/check-in kunjungan."),
 ("Visit", "Approve / reject visit (waiting review)",
  "Buka halaman visits (via dalam leads), approve/reject 1 visit.",
  "Status visit berubah; bulk approve/reject & single check-in berfungsi."),
 ("Settings", "Ganti preset pipeline",
  "Settings > Pipeline Leads: pilih preset (Default/ClientA/Simple), Save & Apply.",
  "Stage berubah; lead ter-remap by-position ke stage baru; board mencerminkan preset."),
 ("Data", "Isolasi localStorage clone",
  "Buka clone lalu induk di tab beda; ubah data di salah satu.",
  "Data clone (satui_proto_data_saclone) terpisah dari induk; tidak saling menimpa."),
]


def write_md():
    lines = [f"# {META['Judul']}", ""]
    for k, v in META.items():
        if k == "Judul":
            continue
        lines.append(f"- **{k}:** {v}")
    lines += ["", "> Kolom Pass/Fail & Catatan diisi tester saat eksekusi. Jangan diisi saat authoring.", ""]
    lines.append("| " + " | ".join(COLS) + " |")
    lines.append("|" + "|".join(["---"] * len(COLS)) + "|")
    for i, (area, sc, step, exp) in enumerate(CASES, 1):
        row = [str(i), area, sc, step, exp, "", ""]
        row = [c.replace("|", "\\|") for c in row]
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    (OUT / "uat-prototype-sales.md").write_text("\n".join(lines), encoding="utf-8")


def write_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UAT Sales"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    r = 1
    ws.cell(r, 1, META["Judul"]).font = Font(bold=True, size=14)
    r += 2
    for k in ["Apps Name", "Version", "Date", "PM", "QA Name", "TechLead", "Scope"]:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, META[k])
        r += 1
    r += 1
    hdr = r
    for c, name in enumerate(COLS, 1):
        cell = ws.cell(r, c, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    r += 1
    for i, (area, sc, step, exp) in enumerate(CASES, 1):
        vals = [i, area, sc, step, exp, "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        r += 1
    widths = [5, 14, 30, 40, 48, 10, 20]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    wb.save(OUT / "uat-prototype-sales.xlsx")


if __name__ == "__main__":
    assert len({sc for _, sc, _, _ in CASES}) == len(CASES), "skenario duplikat"
    write_md()
    write_xlsx()
    print(f"OK {len(CASES)} skenario -> uat/uat-prototype-sales.md + .xlsx")
